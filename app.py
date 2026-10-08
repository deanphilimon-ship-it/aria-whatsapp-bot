from flask import Flask, request
import requests, os, base64, json, re, hashlib, hmac, atexit, threading, time, io, csv, uuid, tempfile
from collections import OrderedDict, defaultdict
from datetime import datetime, timedelta
import pytz
from groq import Groq

# ============================================================
# ARIA - Advanced Responsive Intelligent Assistant
# VERSION 15.2
# ============================================================
app = Flask(__name__)
VERSION = "v15.2"
start_time = time.time()
GRAPH_API_VERSION = os.getenv("GRAPH_API_VERSION", "v26.0")

# ============================================================
# GLOBAL STATE
# ============================================================
last_explain_topic = {}
last_explain_fields = {}
user_waiting_image = {}
IMAGE_WAIT_TIMEOUT = 30 * 60
conversation_memory = {}
user_profile = {}
MAX_HISTORY = 40
RECENT_MEMORY_MESSAGES = 20
LONG_TERM_FACT_LIMIT = 50
MEMORY_FILE = "aria_memory.json"
DOCUMENT_CONTEXT_LIMIT = 12000
document_context = {}
DOCUMENT_WAIT_TIMEOUT = 2 * 60 * 60
MAX_MEMORY_CONTEXT_CHARS = 3500
MAX_PROMPT_CHARS = 14000
MAX_REMINDERS_PER_USER = 20

# ============================================================
# ENVIRONMENT
# ============================================================
WHATSAPP_TOKEN = os.getenv("WHATSAPP_TOKEN")
PHONE_NUMBER_ID = os.getenv("PHONE_NUMBER_ID")
VERIFY_TOKEN = os.getenv("VERIFY_TOKEN")
APP_SECRET = os.getenv("WHATSAPP_APP_SECRET", "").strip()
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
VOICE_MODE = os.getenv("VOICE_MODE", "text").strip().lower()
STT_MODEL = os.getenv("STT_MODEL", "whisper-large-v3-turbo")
TTS_MODEL = os.getenv("TTS_MODEL", "canopylabs/orpheus-v1-english")
TTS_VOICE = os.getenv("TTS_VOICE", "autumn")
TTS_LANGUAGE = os.getenv("TTS_LANGUAGE", "en")
UNSPLASH_KEY = os.getenv("UNSPLASH_KEY")
PIXABAY_KEY = os.getenv("PIXABAY_KEY")
COMICVINE_KEY = os.getenv("COMICVINE_KEY")
PINTEREST_ACCESS_TOKEN = os.getenv("PINTEREST_ACCESS_TOKEN") or os.getenv("PINTEREST_API_KEY") or os.getenv("PINTEREST_TOKEN")
IMAGE_MODEL = os.getenv("IMAGE_MODEL", "flux").strip() or "flux"
IMAGE_WIDTH = int(os.getenv("IMAGE_WIDTH", "1024"))
IMAGE_HEIGHT = int(os.getenv("IMAGE_HEIGHT", "1024"))
IMAGE_ENHANCE = os.getenv("IMAGE_ENHANCE", "true").strip().lower() == "true"
CHAT_MODEL = os.getenv("CHAT_MODEL", "openai/gpt-oss-120b")
VISION_MODEL = os.getenv("VISION_MODEL", "qwen/qwen3.8-27b")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
GEMINI_THINKING_LEVEL = os.getenv("GEMINI_THINKING_LEVEL", "low")
ARIA_TIMEZONE = os.getenv("ARIA_TIMEZONE", "Africa/Lagos")
client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None

# ============================================================
# ACCESS CONTROL / RATE LIMITS
# ============================================================
OWNER_NUMBER = os.getenv("OWNER_NUMBER", "").strip()
OWNER_NAME = "Philimon Dean"
if not OWNER_NUMBER: print("[WARNING] OWNER_NUMBER is not set. Owner commands are disabled.")
ADMIN_NUMBERS = {n.strip() for n in os.getenv("ADMIN_NUMBERS", "").split(",") if n.strip()}
ALLOWED_USERS = [n.strip() for n in os.getenv("ALLOWED_USERS", "").split(",") if n.strip()]
ARIA_PASSWORD = os.getenv("ARIA_PASSWORD", "").strip()
PASSWORD = ARIA_PASSWORD or None
if not PASSWORD: print("[WARNING] ARIA_PASSWORD is not set. Password login is disabled.")
MAX_TRIES = 4
auth_tries, voice_enabled_users, authenticated_users, BANNED_USERS = {}, set(), set(), set()
api_requests = defaultdict(list)
_rate_lock = threading.Lock()
LIMITS = {"pint1":20,"pint2":100,"pint4":20,"pint5":50,"pinterest":20,"create":int(os.getenv("CREATE_LIMIT_PER_HOUR","15")),"live":int(os.getenv("LIVE_LIMIT_PER_HOUR","20")),"ai":int(os.getenv("AI_LIMIT_PER_HOUR","40"))}

def check_rate_limit(number, api_name):
    now=time.time(); key=f"{number}_{api_name}"
    with _rate_lock:
        api_requests[key]=[t for t in api_requests[key] if now-t<3600]
        if number not in ({OWNER_NUMBER}|ADMIN_NUMBERS) and len(api_requests[key])>=LIMITS.get(api_name,100): return False
        api_requests[key].append(now)
    return True

# ============================================================
# PERSISTENT MEMORY
# ============================================================
UPSTASH_URL=os.getenv("UPSTASH_REDIS_REST_URL","").strip().rstrip("/")
UPSTASH_TOKEN=os.getenv("UPSTASH_REDIS_REST_TOKEN","").strip()
REMOTE_ENABLED=bool(UPSTASH_URL and UPSTASH_TOKEN)
KEY_DATA="aria:data"; KEY_BANNED="aria:banned"; REMOTE_FLUSH_SECONDS=int(os.getenv("REMOTE_FLUSH_SECONDS","30"))
_memory_lock=threading.Lock(); _remote_lock=threading.Lock(); _dirty=False; _last_banned=None

def remote_cmd(*args):
    r=requests.post(UPSTASH_URL,headers={"Authorization":f"Bearer {UPSTASH_TOKEN}"},json=list(args),timeout=5); r.raise_for_status(); return r.json().get("result")

def _snapshot():
    for _ in range(3):
        try: return json.dumps({"convo":conversation_memory,"profile":user_profile},ensure_ascii=False)
        except RuntimeError: time.sleep(.05)
    return None

def flush_remote(bans_only=False):
    global _dirty,_last_banned
    if not REMOTE_ENABLED:return
    with _remote_lock:
        try:
            banned_now=sorted(BANNED_USERS)
            if banned_now!=_last_banned:
                remote_cmd("SET",KEY_BANNED,json.dumps(banned_now)); _last_banned=banned_now
            if not bans_only:
                snap=_snapshot()
                if snap: remote_cmd("SET",KEY_DATA,snap); _dirty=False
        except Exception as e: print("[REMOTE SAVE ERROR]",repr(e))

def _apply_memory(data,banned):
    global conversation_memory,user_profile,BANNED_USERS
    conversation_memory=data.get("convo",{}); user_profile=data.get("profile",{}); BANNED_USERS=set(banned)

def load_memory():
    global _last_banned
    data=banned=None
    if REMOTE_ENABLED:
        try:
            raw=remote_cmd("GET",KEY_DATA); raw_b=remote_cmd("GET",KEY_BANNED)
            if raw:data=json.loads(raw)
            if raw_b:banned=json.loads(raw_b); _last_banned=sorted(banned)
            print("[MEMORY] Loaded from Upstash")
        except Exception as e: print("[REMOTE LOAD ERROR]",repr(e))
    if (data is None or banned is None) and os.path.exists(MEMORY_FILE):
        try:
            with open(MEMORY_FILE,encoding="utf-8") as f: local=json.load(f)
            if data is None:data={"convo":local.get("convo",{}),"profile":local.get("profile",{})}
            if banned is None:banned=local.get("banned",[])
            print("[MEMORY] Loaded from local file")
        except Exception as e: print("[MEMORY LOAD ERROR]",repr(e))
    _apply_memory(data or {},banned or [])

def save_memory():
    global _dirty
    try:
        with _memory_lock:
            tmp=MEMORY_FILE+".tmp"
            with open(tmp,"w",encoding="utf-8") as f: json.dump({"convo":conversation_memory,"profile":user_profile,"banned":list(BANNED_USERS)},f,ensure_ascii=False,indent=2)
            os.replace(tmp,MEMORY_FILE)
    except Exception as e: print("[MEMORY SAVE ERROR]",repr(e))
    if REMOTE_ENABLED:
        _dirty=True
        if sorted(BANNED_USERS)!=_last_banned: flush_remote(bans_only=True)

def _remote_flusher():
    while True:
        time.sleep(REMOTE_FLUSH_SECONDS)
        if _dirty: flush_remote()
load_memory()
if REMOTE_ENABLED:
    threading.Thread(target=_remote_flusher,daemon=True).start(); atexit.register(lambda:flush_remote() if _dirty else None); _dirty=True
else: print("[WARNING] Upstash is not configured. Memory and bans only live in the local file and are lost on redeploy.")

# ============================================================
# AUTH
# ============================================================
def check_auth(from_number,text):
    if OWNER_NUMBER and from_number==OWNER_NUMBER:
        BANNED_USERS.discard(from_number); authenticated_users.add(from_number); return True,""
    if from_number in BANNED_USERS:return False,"You are banned from using ARIA."
    if from_number in ADMIN_NUMBERS or from_number in ALLOWED_USERS or from_number in authenticated_users:
        authenticated_users.add(from_number); return True,""
    if PASSWORD and hmac.compare_digest(text.strip().upper().encode(),PASSWORD.upper().encode()):
        authenticated_users.add(from_number); auth_tries[from_number]=0; return True,"Access Granted.\n\nWelcome to ARIA."
    if not text.strip():return False,"Private Bot\n\nSend password to unlock."
    auth_tries[from_number]=auth_tries.get(from_number,0)+1
    if auth_tries[from_number]>=MAX_TRIES:
        BANNED_USERS.add(from_number); save_memory(); return False,"Locked. Too many incorrect attempts."
    return False,f"Private Bot\n\nSend password to unlock.\nAttempts left: {MAX_TRIES-auth_tries[from_number]}"

def get_users_list():
    users=set(authenticated_users)|set(conversation_memory)|set(user_profile)|set(ALLOWED_USERS); users.discard(OWNER_NUMBER)
    if not users:return "〔 *ARIA USERS* 〕\n\nNo users recorded."
    lines=["〔 *ARIA USERS* 〕","",f"*Authenticated:* {len(authenticated_users)}",f"*Recorded:* {len(users)}",f"*Banned:* {len(BANNED_USERS)}",""]
    lines += [f"• {n} — {'BANNED' if n in BANNED_USERS else 'ACTIVE'}" for n in sorted(users)]
    return "\n".join(lines)

# ============================================================
# WHATSAPP OUTPUT
# ============================================================
def clean_ui(text):
    if not text:return ""
    lines=[]
    for line in str(text).splitlines():
        s=line.strip()
        if s.startswith("|") and s.endswith("|"):continue
        if re.match(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?$",s):continue
        lines.append(line)
    text="\n".join(lines)
    text=re.sub(r"^#{1,6}\s*(.+)$",r"〔 *\1* 〕",text,flags=re.MULTILINE).replace("**","*")
    return re.sub(r"\n{4,}","\n\n",text).strip()

def graph_url(endpoint):return f"https://graph.facebook.com/{GRAPH_API_VERSION}/{endpoint}"
def _split_whatsapp_text(text,max_len=3500):
    text=(text or "").strip()
    if len(text)<=max_len:return [text] if text else []
    parts=[]; remaining=text
    while len(remaining)>max_len:
        window=remaining[:max_len]; cut=max(window.rfind("\n\n"),window.rfind("\n"),window.rfind(". "),window.rfind("! "),window.rfind("? "))
        if cut<max_len*.55:cut=max_len
        parts.append(remaining[:cut].strip()); remaining=remaining[cut:].strip()
    if remaining:parts.append(remaining)
    return [x for x in parts if x]

def send_text(to,text):
    text=clean_ui(text) or "I couldn't generate a response."
    if not WHATSAPP_TOKEN or not PHONE_NUMBER_ID:return False
    url=graph_url(f"{PHONE_NUMBER_ID}/messages"); headers={"Authorization":f"Bearer {WHATSAPP_TOKEN}","Content-Type":"application/json"}; chunks=_split_whatsapp_text(text); success=True
    for i,chunk in enumerate(chunks):
        body=(f"〔 ARIA · {i+1}/{len(chunks)} 〕\n\n" if len(chunks)>1 else "")+chunk
        try:
            r=requests.post(url,headers=headers,json={"messaging_product":"whatsapp","to":to,"type":"text","text":{"preview_url":False,"body":body}},timeout=20)
            if not r.ok:success=False; print("[WHATSAPP SEND ERROR]",r.status_code,r.text[:1000])
        except Exception as e:success=False; print("[WHATSAPP SEND EXCEPTION]",repr(e))
        time.sleep(.3)
    return success

def _guess_image_mime(content_type,image_url):
    mime=(content_type or "").split(";",1)[0].strip().lower()
    if mime.startswith("image/"):return mime
    lower=(image_url or "").lower().split("?",1)[0]
    return "image/png" if lower.endswith(".png") else "image/jpeg"

def _upload_whatsapp_media(image_bytes,mime_type):
    r=requests.post(graph_url(f"{PHONE_NUMBER_ID}/media"),headers={"Authorization":f"Bearer {WHATSAPP_TOKEN}"},data={"messaging_product":"whatsapp","type":mime_type},files={"file":("aria_image",image_bytes,mime_type)},timeout=45); r.raise_for_status(); mid=r.json().get("id")
    if not mid:raise RuntimeError("WhatsApp media upload returned no media ID")
    return mid

def send_image_url(to,image_url,caption=""):
    if not image_url:return False
    if not WHATSAPP_TOKEN or not PHONE_NUMBER_ID:return False
    message_url=graph_url(f"{PHONE_NUMBER_ID}/messages"); headers={"Authorization":f"Bearer {WHATSAPP_TOKEN}","Content-Type":"application/json"}
    try:
        r=requests.get(image_url,headers={"User-Agent":"ARIA-Bot/1.0"},timeout=45); r.raise_for_status(); data=r.content
        if not data or len(data)>5*1024*1024:raise ValueError("Image is empty or larger than WhatsApp's 5 MB limit")
        mime=_guess_image_mime(r.headers.get("Content-Type"),image_url)
        if mime not in {"image/jpeg","image/png"}:raise ValueError(f"Unsupported image type: {mime}")
        mid=_upload_whatsapp_media(data,mime)
        s=requests.post(message_url,headers=headers,json={"messaging_product":"whatsapp","to":to,"type":"image","image":{"id":mid,"caption":caption[:1024]}},timeout=30)
        if s.ok:return True
        raise RuntimeError(s.text[:1000])
    except Exception as e:
        print("[WHATSAPP IMAGE ERROR]",repr(e))
        try:
            s=requests.post(message_url,headers=headers,json={"messaging_product":"whatsapp","to":to,"type":"image","image":{"link":image_url,"caption":caption[:1024]}},timeout=30)
            if s.ok:return True
        except Exception as e2:print("[WHATSAPP IMAGE FALLBACK ERROR]",repr(e2))
        send_text(to,"The image could not be delivered."); return False

# ============================================================
# MEDIA / PINT
# ============================================================
def pint1_unsplash(sender,query):
    if not UNSPLASH_KEY:return send_text(sender,"Unsplash is not configured.")
    if not check_rate_limit(sender,"pint1"):return send_text(sender,"Too many Unsplash searches. Try again later.")
    send_text(sender,f"Pint1 Unsplash: *{query}*...")
    try:
        r=requests.get("https://api.unsplash.com/photos/random",params={"query":query,"client_id":UNSPLASH_KEY,"orientation":"portrait"},timeout=15); r.raise_for_status(); x=r.json(); send_image_url(sender,x["urls"]["regular"],f"Pint1: {query}\nPhoto by {x['user']['name']} on Unsplash")
    except Exception as e:print("[UNSPLASH ERROR]",repr(e)); send_text(sender,f"No results on Unsplash for: {query}")

def pint2_pexels(sender,query):
    if not check_rate_limit(sender,"pint2"):return send_text(sender,"Too many wallpaper searches. Try again later.")
    send_text(sender,f"Pint2 Wallhaven: *{query}*...")
    try:
        r=requests.get("https://wallhaven.cc/api/v1/search",params={"q":query,"sorting":"random","atleast":"1920x1080","ratios":"16x9,9x16"},timeout=15); r.raise_for_status(); d=r.json().get("data",[])
        if d and d[0].get("path"):return send_image_url(sender,d[0]["path"],f"Pint2: {query}\n{d[0].get('resolution','Unknown resolution')} | Wallhaven")
        raise ValueError("No wallpapers found")
    except Exception as e:print("[WALLHAVEN ERROR]",repr(e))
    if UNSPLASH_KEY:
        try:
            r=requests.get("https://api.unsplash.com/photos/random",params={"query":query,"client_id":UNSPLASH_KEY,"orientation":"portrait"},timeout=15); r.raise_for_status(); return send_image_url(sender,r.json()["urls"]["regular"],f"Pint2 Fallback: {query}\nFrom Unsplash")
        except Exception as e:print("[UNSPLASH FALLBACK ERROR]",repr(e))
    p=requests.utils.quote(f"{query}, fantasy art, aesthetic wallpaper, ultra detailed, high quality")
    return send_image_url(sender,f"https://image.pollinations.ai/prompt/{p}?model={requests.utils.quote(IMAGE_MODEL)}&width={IMAGE_WIDTH}&height={max(IMAGE_HEIGHT,1536)}&enhance={str(IMAGE_ENHANCE).lower()}&nologo=true",f"Pint2 AI: {query}\nGenerated with FLUX")

def pint3_anime(sender,query):
    q=query.lower().replace(" ","_")
    if any(x in q for x in ["naked","nude","boobs","pussy","sex","nsfw"]):return send_text(sender,"SFW only. Try Goku, Naruto, Luffy, etc.")
    send_text(sender,f"Pint3 Danbooru: *{query}*...")
    try:
        r=requests.get("https://danbooru.donmai.us/posts.json",params={"tags":f"{q} rating:g","limit":1,"random":"true"},headers={"User-Agent":"ARIA-Bot/1.0"},timeout=15); r.raise_for_status(); posts=r.json()
        if not posts:return send_text(sender,f"No results for: *{query}*")
        p=posts[0]; u=p.get("large_file_url") or p.get("file_url"); u=("https://danbooru.donmai.us"+u if u and u.startswith("/") else u)
        if not u:return send_text(sender,"Danbooru returned a result without an image.")
        return send_image_url(sender,u,f"Pint3: {query}\nTags: {', '.join(p.get('tag_string','').split()[:5])}\nSource: Danbooru")
    except Exception as e:print("[DANBOORU ERROR]",repr(e)); send_text(sender,"Danbooru failed. Try a different character.")

def pint4_comics(sender,query):
    if not COMICVINE_KEY:return send_text(sender,"ComicVine is not configured. Add COMICVINE_KEY to Render.")
    if not check_rate_limit(sender,"pint4"):return send_text(sender,"Too many comic searches. Try again later.")
    send_text(sender,f"Pint4 Comics: *{query}*...")
    try:
        r=requests.get("https://comicvine.gamespot.com/api/search/",params={"api_key":COMICVINE_KEY,"format":"json","query":query,"resources":"character,issue,volume","limit":1},headers={"User-Agent":"ARIA-Bot/1.0"},timeout=15); r.raise_for_status(); d=r.json()
        if d.get("status_code")!=1:return send_text(sender,"ComicVine returned an API error.")
        results=d.get("results",[])
        if not results:return send_text(sender,f"No comics found for: *{query}*")
        x=results[0]; name=x.get("name") or x.get("volume",{}).get("name",query); desc=(x.get("deck") or "No description")[:200]; u=(x.get("image") or {}).get("super_url")
        if u:return send_image_url(sender,u,f"Pint4: {name}\n{desc}...\nType: {x.get('resource_type','unknown')}")
        send_text(sender,f"Pint4: {name}\n{desc}...")
    except Exception as e:print("[COMICVINE ERROR]",repr(e)); send_text(sender,"ComicVine failed. Try again in a few seconds.")

def _send_educational_result(sender,result):
    u=result.get("url") or result.get("thumbnail")
    if not u:return False
    c=f"🧪 Education: {(result.get('title') or 'Educational image')[:180]}\nSource: {result.get('source','Openverse')}"
    if result.get("creator"):c+=f"\nCreator: {str(result['creator'])[:120]}"
    if result.get("license"):c+=f"\nLicense: {str(result['license'])[:120]}"
    return send_image_url(sender,u,c)

def _search_wikimedia_education(query):
    r=requests.get("https://commons.wikimedia.org/w/api.php",params={"action":"query","generator":"search","gsrsearch":f"{query} diagram scientific educational apparatus","gsrnamespace":"6","gsrlimit":"8","prop":"imageinfo","iiprop":"url|extmetadata","iiurlwidth":"1400","format":"json","formatversion":"2"},headers={"User-Agent":"ARIA-Bot/1.0"},timeout=15); r.raise_for_status()
    for p in r.json().get("query",{}).get("pages",[]):
        infos=p.get("imageinfo") or []
        if not infos:continue
        info=infos[0]
        if not (info.get("mime") or "").lower().startswith("image/"):continue
        ext=info.get("extmetadata") or {}
        def mv(k):
            v=ext.get(k,{}).get("value",""); return re.sub(r"<[^>]+>","",v).strip() if isinstance(v,str) else ""
        return {"url":info.get("thumburl") or info.get("url"),"title":p.get("title","Educational image").replace("File:",""),"source":"Wikimedia Commons","creator":mv("Artist"),"license":mv("LicenseShortName")}

def _search_openverse_education(query):
    r=requests.get("https://api.openverse.org/v1/images/",params={"q":f"{query} educational diagram scientific apparatus","page_size":8},headers={"User-Agent":"ARIA-Bot/1.0"},timeout=15); r.raise_for_status()
    for x in r.json().get("results",[]):
        u=x.get("url") or x.get("thumbnail")
        if u:return {"url":u,"thumbnail":x.get("thumbnail"),"title":x.get("title") or "Educational image","source":"Openverse","creator":x.get("creator"),"license":x.get("license")}

def pint5_education(sender,query):
    if not check_rate_limit(sender,"pint5"):return send_text(sender,"Too many education searches. Try again later.")
    send_text(sender,f"🧪 Education: *{query}*\nSearching educational diagrams...")
    for fn in (_search_wikimedia_education,_search_openverse_education):
        try:
            x=fn(query)
            if x and _send_educational_result(sender,x):return
        except Exception as e:print("[PINT5 ERROR]",repr(e))
    send_text(sender,f"I couldn't find a suitable educational diagram for *{query}*.\n\nI did not fall back to Unsplash because `.edu` is reserved for educational material.")

def pint_pinterest(sender,query):
    if not PINTEREST_ACCESS_TOKEN:return send_text(sender,"Pinterest is not configured.\n\nAdd `PINTEREST_ACCESS_TOKEN` to Render.")
    if not check_rate_limit(sender,"pinterest"):return send_text(sender,"Too many Pinterest searches. Try again later.")
    send_text(sender,f"📌 Pinterest: *{query}*...")
    try:
        r=requests.get("https://api.pinterest.com/v5/search/pins",headers={"Authorization":f"Bearer {PINTEREST_ACCESS_TOKEN}","Content-Type":"application/json","User-Agent":"ARIA-Bot/1.0"},params={"query":query,"page_size":25},timeout=20)
        if r.status_code==401:return send_text(sender,"Pinterest rejected the access token (401). Generate a fresh production token with `pins:read` and update `PINTEREST_ACCESS_TOKEN` in Render.")
        if r.status_code==403:return send_text(sender,"Pinterest refused the search request (403). Your app/token may not have access to Pin search yet.")
        r.raise_for_status(); items=r.json().get("items",[])
        if not items:return send_text(sender,f"No Pinterest Pins matched *{query}*.")
        pin=next((x for x in items if (x.get("media") or {}).get("images")),items[0]); images=(pin.get("media") or {}).get("images") or {}; u=None
        for k in ("1200x","600x","400x300","150x150","orig","originals"):
            v=images.get(k); u=v.get("url") if isinstance(v,dict) else v
            if u:break
        if not u:return send_text(sender,"Pinterest returned a matching Pin without a usable image URL.")
        c=f"📌 Pinterest: {pin.get('title') or query}\nSource: Pinterest"+(f"\n{pin['link']}" if pin.get("link") else "")
        send_image_url(sender,u,c)
    except Exception as e:print("[PINTEREST EXCEPTION]",repr(e)); send_text(sender,"Pinterest search failed. Check the token/scopes and try again.")

def imagine_generate(sender,prompt):
    send_text(sender,f"FLUX AI generating:\n*{prompt}*...")
    enhanced=prompt+", ultra detailed, cinematic lighting, sharp focus, high quality, professional artwork"; encoded=requests.utils.quote(enhanced); seed=uuid.uuid4().int%100000
    url=f"https://image.pollinations.ai/prompt/{encoded}?model={requests.utils.quote(IMAGE_MODEL)}&width={IMAGE_WIDTH}&height={IMAGE_HEIGHT}&enhance={str(IMAGE_ENHANCE).lower()}&nologo=true&seed={seed}"
    if not send_image_url(sender,url,f"FLUX: {prompt}\nModel: FLUX"):
        turbo=f"https://image.pollinations.ai/prompt/{encoded}?model=turbo&width=1024&height=1024&nologo=true"; send_image_url(sender,turbo,f"AI: {prompt}\nFallback: Turbo")

# ============================================================
# AI / MEMORY
# ============================================================
DEFAULT_SYSTEM_PROMPT="""You are ARIA, an advanced WhatsApp AI assistant.
IDENTITY
- Your name is ARIA.
- ARIA stands for Advanced Responsive Intelligent Assistant.
- The owner of this ARIA instance is Philimon Dean.
- Be intelligent, practical, direct and conversational.
- Do not pretend to be human or claim access you do not have.
RESPONSE STYLE
- Respond naturally for WhatsApp. Keep normal answers concise.
- Use short paragraphs and bullets. No markdown tables or ### headings.
- Do not add filler or unnecessarily repeat the question.
ACCURACY
- If uncertain, say so. Never invent sources, facts, API results, files, events or memories.
- If current information is requested without live data, state the limitation.
MEMORY
- Use supplied memory when relevant. Do not claim to remember absent information.
SAFETY
- Do not generate explicit sexual content. Educational, medical, anatomical, safety and non-explicit discussions may be handled clinically.
"""

def build_memory_context(from_number,current_input="",skip_last_user=False,max_chars=None):
    max_chars=max_chars or MAX_MEMORY_CONTEXT_CHARS
    p=user_profile.get(from_number,{}); head=[]
    if p.get("name"):head.append(f"You are talking to {p['name']}.")
    facts=p.get("facts",[]); prefs=p.get("preferences",[])
    if facts:head.append("Long-term facts:\n"+"\n".join(f"- {str(x)[:200]}" for x in facts[-15:]))
    if prefs:head.append("Preferences:\n"+"\n".join(f"- {str(x)[:200]}" for x in prefs[-10:]))
    head_text="\n\n".join(head)
    msgs=list(conversation_memory.get(from_number,[]))
    if skip_last_user and msgs and msgs[-1].get("role")=="user":msgs=msgs[:-1]
    budget=max(max_chars-len(head_text),800)
    older_block=""
    if len(msgs)>RECENT_MEMORY_MESSAGES and current_input:
        stop={"what","when","where","which","that","this","with","from","about"}
        terms={x.lower() for x in re.findall(r"[a-zA-Z0-9]{4,}",current_input) if x.lower() not in stop}
        scored=[]
        for m in msgs[:-RECENT_MEMORY_MESSAGES]:
            c=str(m.get("content","")); score=sum(c.lower().count(t) for t in terms)
            if score:scored.append((score,c,m.get("role","unknown")))
        scored.sort(key=lambda x:x[0],reverse=True)
        if scored:older_block=("Relevant older memory:\n"+"\n".join(f"{r}: {c[:300]}" for _,c,r in scored[:4]))[:budget//3]
    recent_budget=budget-len(older_block); kept=[]; used=0
    for m in reversed(msgs[-RECENT_MEMORY_MESSAGES:]):
        line=f"{m.get('role','unknown')}: {str(m.get('content',''))[:400]}"
        if kept and used+len(line)+1>recent_budget:break
        kept.append(line); used+=len(line)+1
    kept.reverse()
    out=[]
    if head_text:out.append(head_text)
    if kept:out.append("Recent conversation:\n"+"\n".join(kept))
    if older_block:out.append(older_block)
    return "\n\n".join(out)

_NO_MEMORY_PREFIXES=("AI request failed","AI is not configured","You've reached the hourly","I couldn't","No recent image","No recent document")
def add_to_memory(from_number,role,content):
    if not str(content or "").strip():return
    if role=="assistant" and str(content).startswith(_NO_MEMORY_PREFIXES):return
    conversation_memory.setdefault(from_number,[]).append({"role":role,"content":content}); conversation_memory[from_number]=conversation_memory[from_number][-MAX_HISTORY:]; save_memory()

def remember_user_fact(from_number,fact):
    fact=re.sub(r"\s+"," ",fact).strip().rstrip(".")
    if not fact:return None
    p=user_profile.setdefault(from_number,{}); facts=p.setdefault("facts",[])
    if fact not in facts:facts.append(fact); p["facts"]=facts[-LONG_TERM_FACT_LIMIT:]; save_memory()
    return f"Got it. I'll remember that: *{fact}*."

def forget_user_memory(from_number):
    conversation_memory.pop(from_number,None); user_profile.pop(from_number,None); document_context.pop(from_number,None); user_waiting_image.pop(from_number,None); save_memory()

def learn_fact(from_number,text):
    m=re.search(r"my name is\s+(.+)",text,re.I)
    if m and "remember" in text.lower():
        name=m.group(1).strip(); user_profile.setdefault(from_number,{})["name"]=name; save_memory(); return f"Got it. I'll remember your name is *{name}*."
    m=re.match(r"(?:remember that|remember)\s+(.+)$",text,re.I)
    if m:return remember_user_fact(from_number,m.group(1))

def _extract_groq_text(c):
    if not c or not getattr(c,"choices",None):return ""
    msg=getattr(c.choices[0],"message",None); content=getattr(msg,"content",None) if msg else None
    return content.strip() if isinstance(content,str) else str(content).strip() if content is not None else ""

def _extract_gemini_text(data):
    texts=[]
    for part in ((data.get("candidates") or [{}])[0].get("content") or {}).get("parts") or []:
        if part.get("text"):texts.append(str(part["text"]))
    return "\n".join(texts).strip()
GEMINI_API_URL="https://generativelanguage.googleapis.com/v1beta/models/"

def gemini_call(prompt,system=DEFAULT_SYSTEM_PROMPT,image_data=None,mime_type=None,timeout=45):
    if not GEMINI_API_KEY:raise RuntimeError("GEMINI_API_KEY is missing")
    parts=[]
    if image_data:parts.append({"inline_data":{"mime_type":mime_type,"data":image_data}})
    parts.append({"text":prompt}); payload={"system_instruction":{"parts":[{"text":system}]},"contents":[{"role":"user","parts":parts}],"generationConfig":{"maxOutputTokens":1024,"thinkingConfig":{"thinkingLevel":GEMINI_THINKING_LEVEL}}}
    r=requests.post(f"{GEMINI_API_URL}{GEMINI_MODEL}:generateContent",headers={"Content-Type":"application/json","x-goog-api-key":GEMINI_API_KEY},json=payload,timeout=timeout); r.raise_for_status(); text=_extract_gemini_text(r.json())
    if not text:raise ValueError("Gemini returned an empty response")
    return text

def groq_chat_call(prompt,system=DEFAULT_SYSTEM_PROMPT):
    if not client:raise RuntimeError("GROQ_API_KEY is missing")
    c=client.chat.completions.create(model=CHAT_MODEL,messages=[{"role":"system","content":system},{"role":"user","content":prompt}],temperature=.6,max_tokens=1024,top_p=.95,stream=False); text=_extract_groq_text(c)
    if not text:raise ValueError("Groq returned an empty response")
    return text

def ai_call(prompt,from_number,system=DEFAULT_SYSTEM_PROMPT):
    if from_number not in ({OWNER_NUMBER}|ADMIN_NUMBERS) and not check_rate_limit(from_number,"ai"):return "You've reached the hourly AI limit. Please try again in a while."
    ctx=build_memory_context(from_number,prompt,skip_last_user=True); full=f"Current date/time: {_now_str()}\n\n"+(ctx+"\n\n" if ctx else "")+f"User: {str(prompt)[:MAX_PROMPT_CHARS]}"; errors=[]
    if client:
        try:return groq_chat_call(full,system)
        except Exception as e:errors.append(f"Groq: {e}")
    if GEMINI_API_KEY:
        try:return gemini_call(full,system)
        except Exception as e:errors.append(f"Gemini: {e}")
    print("[AI ERROR SUMMARY]"," | ".join(map(str,errors))); return "AI request failed.\n\nBoth configured AI providers failed. Use `.health` to see their status." if (client or GEMINI_API_KEY) else "AI is not configured.\n\nAdd GROQ_API_KEY or GEMINI_API_KEY to Render."

# ============================================================
# VISION / VOICE
# ============================================================
def download_whatsapp_image(image_ref,mime_type=None):
    if not image_ref or not WHATSAPP_TOKEN:raise RuntimeError("WhatsApp image credentials/media are missing")
    h={"Authorization":f"Bearer {WHATSAPP_TOKEN}"}
    if str(image_ref).startswith(("http://","https://")):url=image_ref; resolved=mime_type
    else:
        r=requests.get(graph_url(str(image_ref)),headers=h,timeout=20); r.raise_for_status(); info=r.json(); url=info.get("url"); resolved=mime_type or info.get("mime_type")
        if not url:raise ValueError("Meta did not return a media URL.")
    r=requests.get(url,headers=h,timeout=30); r.raise_for_status()
    detected=(resolved or r.headers.get("Content-Type","image/jpeg")).split(";",1)[0].strip().lower()
    if not detected.startswith("image/"):raise ValueError(f"Downloaded media is not an image: {detected}")
    if len(r.content)>5*1024*1024:raise ValueError("Image is larger than WhatsApp's 5 MB image limit.")
    return base64.b64encode(r.content).decode(),detected

def groq_vision_call(base64_image,detected_mime,prompt):
    if not client:raise RuntimeError("GROQ_API_KEY is missing")
    c=client.chat.completions.create(model=VISION_MODEL,messages=[{"role":"user","content":[{"type":"text","text":prompt+"\n\nBe accurate and concise. If something cannot be determined, say so."},{"type":"image_url","image_url":{"url":f"data:{detected_mime};base64,{base64_image}"}}]}],temperature=.3,max_tokens=1024,top_p=.9,stream=False); return _extract_groq_text(c)

def vision_call(image_url,prompt,mime_type=None):
    try:b64,mime=download_whatsapp_image(image_url,mime_type)
    except Exception as e:return f"I couldn't retrieve that image from WhatsApp.\n\nReason: {str(e)[:500]}"
    if client:
        try:
            out=groq_vision_call(b64,mime,prompt).strip()
            if out:return out
            print("[VISION GROQ] empty response, trying fallback")
        except Exception as e:print("[VISION GROQ ERROR]",repr(e))
    if GEMINI_API_KEY:
        try:return gemini_call(prompt+"\n\nBe accurate and concise. If something cannot be determined, say so.",image_data=b64,mime_type=mime).strip()
        except Exception as e:print("[VISION GEMINI ERROR]",repr(e))
    return "I couldn't process that image.\n\nBoth configured vision providers failed. Use `.health` to see their status."

def solve_image_problem(image_url,mime_type=None,mode="universal"):
    prompt="""You are ARIA's dedicated mathematics solver. Read the image carefully and solve the visible mathematical problem. Use 〔 SOLUTION 〕, Given/question, Method or formula, Key working steps, Final answer. If something is unreadable, say so instead of guessing.""" if mode=="math" else """You are ARIA's universal image-problem solver. Identify whether the image contains mathematics, physics, chemistry, biology, English, multiple choice, logic, a diagram or another academic/practical task, then solve it. Use 〔 SOLUTION 〕, Problem type, Question, Method/formula/rule, Important working, then 〔 ANSWER 〕. Do not invent unreadable text."""
    return vision_call(image_url,prompt,mime_type)

def download_whatsapp_audio(media_id):
    if not WHATSAPP_TOKEN or not media_id:raise RuntimeError("WhatsApp audio credentials/media are missing")
    h={"Authorization":f"Bearer {WHATSAPP_TOKEN}"}; r=requests.get(graph_url(media_id),headers=h,timeout=20); r.raise_for_status(); info=r.json(); url=info.get("url"); mime=info.get("mime_type") or "audio/ogg"
    if not url:raise ValueError("Meta did not return an audio media URL")
    r=requests.get(url,headers=h,timeout=30); r.raise_for_status()
    if len(r.content)>25*1024*1024:raise ValueError("Audio exceeds the 25 MB transcription limit")
    return r.content,mime

def transcribe_audio(audio_bytes,mime_type="audio/ogg"):
    if not client:raise RuntimeError("GROQ_API_KEY is missing")
    ext=".ogg" if "ogg" in mime_type else ".mp3" if "mpeg" in mime_type or "mp3" in mime_type else ".wav" if "wav" in mime_type else ".m4a" if "mp4" in mime_type or "m4a" in mime_type else ".webm"
    x=client.audio.transcriptions.create(file=(f"aria_voice{ext}",audio_bytes,mime_type),model=STT_MODEL,language=TTS_LANGUAGE or None,response_format="json",temperature=0.0); return (getattr(x,"text","") or "").strip()

def generate_tts_audio(text):
    if not GROQ_API_KEY:raise RuntimeError("GROQ_API_KEY is missing")
    clean=re.sub(r"[*_`#]","",text or "").strip(); clean=re.sub(r"\s+"," ",clean)[:3500]
    r=requests.post("https://api.groq.com/openai/v1/audio/speech",headers={"Authorization":f"Bearer {GROQ_API_KEY}","Content-Type":"application/json"},json={"model":TTS_MODEL,"voice":TTS_VOICE,"input":clean,"response_format":"ogg","speed":1.0},timeout=60); r.raise_for_status(); return r.content

def send_audio_bytes(to,audio_bytes,mime_type="audio/ogg"):
    if not WHATSAPP_TOKEN or not PHONE_NUMBER_ID:return False
    r=requests.post(graph_url(f"{PHONE_NUMBER_ID}/media"),headers={"Authorization":f"Bearer {WHATSAPP_TOKEN}"},data={"messaging_product":"whatsapp","type":mime_type},files={"file":("aria_voice.ogg",audio_bytes,mime_type)},timeout=45)
    if not r.ok:return False
    mid=r.json().get("id");
    if not mid:return False
    s=requests.post(graph_url(f"{PHONE_NUMBER_ID}/messages"),headers={"Authorization":f"Bearer {WHATSAPP_TOKEN}","Content-Type":"application/json"},json={"messaging_product":"whatsapp","to":to,"type":"audio","audio":{"id":mid}},timeout=30); return s.ok

def send_voice_reply(to,text):
    try:return send_audio_bytes(to,generate_tts_audio(text))
    except Exception as e:print("[VOICE REPLY ERROR]",repr(e)); return False

# ============================================================
# DOCUMENT READER
# ============================================================
SUPPORTED_DOCUMENT_TYPES={"pdf","docx","txt","md","csv","xlsx","pptx"}; DOCUMENT_MAX_BYTES=15*1024*1024

def _clean_document_text(s):return re.sub(r"\n{4,}","\n\n",re.sub(r"[ \t]+"," ",str(s or ""))).strip()
def _read_pdf_document(data):
    from pypdf import PdfReader
    r=PdfReader(io.BytesIO(data)); return _clean_document_text("\n\n".join(p.extract_text() or "" for p in r.pages))
def _read_docx_document(data):
    from docx import Document
    d=Document(io.BytesIO(data)); return _clean_document_text("\n".join(p.text for p in d.paragraphs)+"\n"+"\n".join(" | ".join(c.text for c in row.cells) for t in d.tables for row in t.rows))
def _read_text_document(data):return _clean_document_text(data.decode("utf-8",errors="replace"))
def _read_csv_document(data):
    rows=list(csv.reader(io.StringIO(data.decode("utf-8",errors="replace")))); return _clean_document_text("\n".join(" | ".join(r) for r in rows))
def _read_xlsx_document(data):
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True); chunks=[]
    for ws in wb.worksheets:
        chunks.append(f"[Sheet: {ws.title}]\n"+"\n".join(" | ".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)))
    return _clean_document_text("\n\n".join(chunks))
def _pptx_shape_text(shape):
    try:
        if getattr(shape,"shape_type",None)==6:return "\n".join(filter(None,(_pptx_shape_text(s) for s in shape.shapes)))
        if getattr(shape,"has_table",False) and shape.has_table:return "\n".join(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
        if getattr(shape,"has_text_frame",False) and shape.has_text_frame:return shape.text_frame.text
    except Exception:pass
    return ""
def _read_pptx_document(data):
    from pptx import Presentation
    p=Presentation(io.BytesIO(data)); slides=[]
    for i,slide in enumerate(p.slides,1):
        body="\n".join(filter(None,(_pptx_shape_text(s) for s in slide.shapes))); notes=""
        try:
            if slide.has_notes_slide:notes=slide.notes_slide.notes_text_frame.text.strip()
        except Exception:pass
        slides.append(f"[Slide {i}]\n{body}"+(f"\n[Notes] {notes}" if notes else ""))
    return _clean_document_text("\n\n".join(slides))
_MIME_EXT={"application/pdf":"pdf","application/vnd.openxmlformats-officedocument.wordprocessingml.document":"docx","text/plain":"txt","text/markdown":"md","text/csv":"csv","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":"xlsx","application/vnd.openxmlformats-officedocument.presentationml.presentation":"pptx"}
def parse_document(filename,data,mime=None):
    if len(data)>DOCUMENT_MAX_BYTES:raise ValueError("Document exceeds the 15 MB limit.")
    ext=filename.lower().rsplit(".",1)[-1].strip() if "." in filename else _MIME_EXT.get((mime or "").split(";")[0].strip().lower(),"")
    if ext not in SUPPORTED_DOCUMENT_TYPES:raise ValueError("Unsupported document type. Supported: PDF, DOCX, TXT, MD, CSV, XLSX, PPTX.\n\nOld .doc/.xls/.ppt files must be re-saved as .docx/.xlsx/.pptx first.")
    text={"pdf":_read_pdf_document,"docx":_read_docx_document,"txt":_read_text_document,"md":_read_text_document,"csv":_read_csv_document,"xlsx":_read_xlsx_document,"pptx":_read_pptx_document}[ext](data)
    if not text:raise ValueError("No readable text was found. A scanned/image-only PDF may require OCR/vision.")
    return text
def store_document_context(user,filename,text):document_context[user]={"filename":filename,"text":text[:DOCUMENT_CONTEXT_LIMIT],"saved_at":time.time()}
def get_document_context(user):
    d=document_context.get(user)
    if not d:return None
    if time.time()-d.get("saved_at",0)>DOCUMENT_WAIT_TIMEOUT:document_context.pop(user,None);return None
    return d
def ask_about_document(user,question):
    d=get_document_context(user)
    if not d:return "No recent document is available. Send a document first."
    return ai_call(f"Document: {d['filename']}\n\nDocument text:\n{d['text']}\n\nUser question: {question}\nAnswer using the document. If the answer is not supported by the document, say so.",user)

# ============================================================
# REMINDERS
# ============================================================
REMINDER_FILE="aria_reminders.json"; reminders={}; _reminder_lock=threading.Lock()
KEY_REMINDERS="aria:reminders"; _reminders_synced=not REMOTE_ENABLED
def _sync_reminders_remote():
    global _reminders_synced
    if _reminders_synced or not REMOTE_ENABLED:return
    raw=remote_cmd("GET",KEY_REMINDERS)
    if raw:
        for k,v in json.loads(raw).items():reminders.setdefault(k,v)
    _reminders_synced=True
def load_reminders():
    global reminders
    data={}
    try:
        if os.path.exists(REMINDER_FILE):
            with open(REMINDER_FILE,encoding="utf-8") as f:data=json.load(f)
    except Exception as e:print("[REMINDER LOAD ERROR]",repr(e))
    reminders=data if isinstance(data,dict) else {}
    try:
        _sync_reminders_remote()
        if REMOTE_ENABLED:print("[REMINDERS] Synced with Upstash")
    except Exception as e:print("[REMINDER REMOTE LOAD ERROR]",repr(e))
def save_reminders():
    try:
        tmp=REMINDER_FILE+".tmp"
        with open(tmp,"w",encoding="utf-8") as f:json.dump(reminders,f,ensure_ascii=False,indent=2)
        os.replace(tmp,REMINDER_FILE)
    except Exception as e:print("[REMINDER SAVE ERROR]",repr(e))
    if REMOTE_ENABLED:
        try:
            _sync_reminders_remote()
            remote_cmd("SET",KEY_REMINDERS,json.dumps(reminders,ensure_ascii=False))
        except Exception as e:print("[REMINDER REMOTE SAVE ERROR]",repr(e))
def _tz():
    try:return pytz.timezone(ARIA_TIMEZONE)
    except Exception:return pytz.utc
def _parse_clock(s):
    m=re.fullmatch(r"(\d{1,2})(?::(\d{2}))?\s*(am|pm)?",s.strip(),re.I)
    if not m:return None
    h=int(m.group(1)); minute=int(m.group(2) or 0); ap=(m.group(3) or "").lower()
    if minute>59:return None
    if ap:
        if h<1 or h>12:return None
        if ap=="pm" and h!=12:h+=12
        if ap=="am" and h==12:h=0
    elif h>23:return None
    return h,minute

_WEEKDAYS={"monday":0,"tuesday":1,"wednesday":2,"thursday":3,"friday":4,"saturday":5,"sunday":6}
def _clock_token(tok):
    t=tok.strip().lower()
    if t=="noon":return 12,0
    if t=="midnight":return 0,0
    return _parse_clock(t)
def _clean_task(s):
    s=re.sub(r"\s+"," ",s).strip(" ,.;:-")
    s=re.sub(r"^(?:to|that)\s+","",s,flags=re.I)
    s=re.sub(r"\s+(?:to|on|at|by)$","",s,flags=re.I)
    return s.strip(" ,.;:-")[:300]
def _loc(tz,d,h,m):return tz.localize(datetime(d.year,d.month,d.day,h,m))

def parse_reminder(text):
    raw=re.sub(r"\b([ap])\.m\.?",r"\1m",text.strip(),flags=re.I)
    body=re.sub(r"^\.?remind(?:ers?)?\b","",raw,flags=re.I).strip()
    body=re.sub(r"^me\b\s*","",body,flags=re.I).strip()
    tz=_tz(); now=datetime.now(tz)
    # relative: "in 30 minutes", "in 2 hours", "in 3 days" (anywhere in the sentence)
    m=re.search(r"\bin\s+(\d+)\s*(minutes?|mins?|hours?|hrs?|days?)\b",body,re.I)
    if m:
        n=int(m.group(1)); u=m.group(2).lower(); secs=n*(60 if u.startswith("min") else 3600 if u.startswith(("hour","hr")) else 86400)
        if n<=0 or secs>365*86400:return None,None
        task=_clean_task(body[:m.start()]+" "+body[m.end():])
        return (task,now+timedelta(seconds=secs)) if task else (None,None)
    # absolute: optional day word + a clock time
    dm=re.search(r"\b(?:on\s+)?(?:next\s+)?(tomorrow|today|tonight|monday|tuesday|wednesday|thursday|friday|saturday|sunday)\b",body,re.I)
    cms=list(re.finditer(r"\bat\s+(noon|midnight|\d{1,2}(?::\d{2})?\s*(?:am|pm)?)(?!\w)|(?<![\w:])(\d{1,2}(?::\d{2})?\s*(?:am|pm))(?!\w)",body,re.I))
    if not cms:return None,None
    cm=cms[-1]; clock=_clock_token(cm.group(1) or cm.group(2))
    if not clock:return None,None
    h,mi=clock; spans=[cm.span()]+([dm.span()] if dm else []); rest=body
    for a,b in sorted(spans,reverse=True):rest=rest[:a]+" "+rest[b:]
    task=_clean_task(rest)
    if not task:return None,None
    base=now.date(); word=dm.group(1).lower() if dm else None
    if word=="tomorrow":base=base+timedelta(days=1)
    elif word in _WEEKDAYS:base=base+timedelta(days=(_WEEKDAYS[word]-base.weekday())%7)
    dt=_loc(tz,base,h,mi)
    if dt<=now:
        if word in ("today","tonight","tomorrow"):return None,None
        dt=_loc(tz,base+timedelta(days=7 if word in _WEEKDAYS else 1),h,mi)
    return task,dt

def create_reminder(user,task,when):
    rid=uuid.uuid4().hex[:8]
    with _reminder_lock:
        reminders[rid]={"id":rid,"user":user,"task":task,"due":when.isoformat(),"created":time.time()}; save_reminders()
    return rid
def list_reminders(user):
    with _reminder_lock:return sorted((dict(r) for r in reminders.values() if r.get("user")==user),key=lambda r:r.get("due",""))
def cancel_reminder(user,rid):
    with _reminder_lock:
        r=reminders.get(rid)
        if r and r.get("user")==user:reminders.pop(rid,None); save_reminders(); return True
    return False
def _fmt_due(r):
    try:return datetime.fromisoformat(r["due"]).astimezone(_tz()).strftime("%d %b %Y %I:%M %p")
    except Exception:return "unknown time"

def reminder_worker():
    while True:
        time.sleep(10); due=[]; now=datetime.now(pytz.utc)
        with _reminder_lock:
            for rid,r in list(reminders.items()):
                if r.get("retry_at",0)>time.time():continue
                try:
                    dt=datetime.fromisoformat(r["due"]); dt=dt if dt.tzinfo else _tz().localize(dt); dt=dt.astimezone(pytz.utc)
                    if dt<=now:due.append((rid,dict(r),dt))
                except Exception:due.append((rid,dict(r),None))
        for rid,r,dt in due:
            late=dt is not None and (now-dt).total_seconds()>300
            ok=send_text(r.get("user"),f"⏰ *Reminder*{' (late)' if late else ''}\n\n{r.get('task','Reminder')}\n\nID: `{rid}`")
            with _reminder_lock:
                cur=reminders.get(rid)
                if not cur:continue
                if ok or dt is None:reminders.pop(rid,None)
                else:
                    cur["attempts"]=cur.get("attempts",0)+1; cur["retry_at"]=time.time()+60*cur["attempts"]
                    if cur["attempts"]>=8:reminders.pop(rid,None)
                save_reminders()
load_reminders(); threading.Thread(target=reminder_worker,daemon=True).start()

# ============================================================
# EDUCATION / STATUS / SAFETY
# ============================================================
EXPLAIN_FIELDS=["Biology","Chemistry","Pharmacology","Clinical","Pathophysiology","Exam Tips"]
def ai_explain(topic,field_num,from_number):
    fields=EXPLAIN_FIELDS
    if str(field_num)=="all":last_explain_topic[from_number]=topic;return f"〔 *{topic.title()} - 6 FIELDS* 〕\n\n"+"\n".join(f"{i}. {x}" for i,x in enumerate(fields,1))+f"\n\nReply `.explain {topic} 3` for Pharmacology."
    try:idx=int(field_num)-1
    except ValueError:return "Usage: `.explain psychology 2`"
    if idx<0 or idx>=len(fields):return "Choose a field from 1 to 6."
    field=fields[idx]
    return ai_call(f"Deep dive into '{topic}' from the {field} perspective. Use examples and exam-relevant details.",from_number,system="You are a knowledgeable professor helping a student. Explain accurately, teach rather than merely define, use clear examples, no markdown tables, and highlight exam-relevant points.")

def get_youtube_link(query):return f"*{query.title()}*\n\n▶️ Tap to search: https://www.youtube.com/results?search_query={requests.utils.quote(query)}"
def get_runtime():
    s=int(time.time()-start_time);return f"{s//3600}h {(s%3600)//60}m {s%60}s"
EXPLICIT_PATTERNS=[r"\bporn(?:ography)?\b",r"\bpornographic\b",r"\bnsfw\b",r"\bsex video\b",r"\bsex tape\b",r"\bexplicit sex\b",r"\bsexual intercourse\b",r"\bgenital(?:s)?\b.*\bphoto(?:s)?\b",r"\bnude(?:d)?\s+(?:photo|picture|image|pic)s?\b",r"\bnaked\s+(?:photo|picture|image|pic)s?\b"]
def is_explicit_request(text):return bool(text) and any(re.search(p,re.sub(r"\s+"," ",str(text).lower()).strip(),re.I) for p in EXPLICIT_PATTERNS)
def safety_block_message():return "I can help with educational, medical, safety, or non-explicit topics, but I can't generate or provide explicit sexual content."
def get_about():return f"""〔 ABOUT ARIA 〕\n\nARIA stands for Advanced Responsive Intelligent Assistant.\n\nARIA is a WhatsApp AI assistant designed for conversation, vision, problem solving, study help, documents, reminders and media commands.\n\nVersion: {VERSION}\nOwner: Philimon Dean\nOwner protection: ENABLED\nContent safety: ENABLED\nVoice: CONFIGURABLE\n"""
def get_menu():return f"""〔 ARIA {VERSION} 〕\n\nAI\n• .ask <question>\n• .explain <topic>\n• .summarize <text>\n• .translate <language> <text>\n• .define <word>\n\nLIVE WEB (Gemini + Google Search)\n• .search <question>\n• .news <topic>\n• .verify <claim>\n• Time-sensitive questions search the web automatically\n\nVISION\n• .describe / .describe detailed\n• .describe ask <question>\n• .read\n• .solve\n• .math\n• .verify (checks edits/AI signs + live fact-check)\n\nSTUDY\n• .study <topic>\n• .quiz <topic>\n\nDOCUMENTS\n• Send PDF/DOCX/TXT/MD/CSV/XLSX/PPTX\n• Ask questions about the latest document\n• .doc <question>\n\nMEDIA\n• .photo <query>\n• .wallpaper <query>\n• .anime <character>\n• .comic <query>\n• .edu <subject>\n• .pinterest <query>\n• .create <prompt>\n• .imagine <prompt> (legacy alias)\n• .play <song name>\n\nREMINDERS\n• .remind me to study at 8pm\n• .remind me tomorrow at 7am to call John\n• .remind me in 30 minutes to check\n• .reminders\n• .remind cancel <id>\n\nSYSTEM\n• .about\n• .health\n• .voice on/off\n• .status\n• .menu\n• .help <command>\n\nOWNER\n• .users\n• .ban <number>\n• .unban <number>\n\nLegacy .pint1–.pint6 still work."""
def get_status():
    local=datetime.now(_tz()).strftime("%I:%M %p"); return f"""〔 *ARIA STATUS* 〕\n\n*Version:* {VERSION}\n*Runtime:* {get_runtime()}\n*Local Time:* {local}\n*Graph API:* {GRAPH_API_VERSION}\n\n〔 *PROVIDERS* 〕\n*Groq:* {'CONFIGURED' if client else 'NOT CONFIGURED'}\n*Groq Chat:* {CHAT_MODEL}\n*Groq Vision:* {VISION_MODEL}\n*Gemini:* {'CONFIGURED' if GEMINI_API_KEY else 'NOT CONFIGURED'}\n*Gemini Model:* {GEMINI_MODEL}\n*Pinterest:* {'CONFIGURED' if PINTEREST_ACCESS_TOKEN else 'NOT CONFIGURED'}\n*Voice Mode:* {VOICE_MODE}\n*STT:* {STT_MODEL}\n*TTS:* {TTS_MODEL} / {TTS_VOICE}\n\n〔 *SYSTEMS* 〕\n*WhatsApp:* {'CONFIGURED' if WHATSAPP_TOKEN and PHONE_NUMBER_ID else 'NOT CONFIGURED'}\n*Memory:* OK\n*Pending Images:* {len(user_waiting_image)}\n*Pending Documents:* {len(document_context)}\n*Reminders:* {len(reminders)} ({'Upstash' if REMOTE_ENABLED else 'LOCAL ONLY - lost on redeploy'})\n*Live Search:* {'ON (Gemini + Google Search)' if LIVE_ENABLED else 'OFF - needs GEMINI_API_KEY'}\n*Image Fallback:* {'ENABLED' if GEMINI_API_KEY else 'DISABLED'}\n"""
def check_groq_chat_health():
    if not client:return "NOT CONFIGURED"
    try:return "ONLINE" if _extract_groq_text(client.chat.completions.create(model=CHAT_MODEL,messages=[{"role":"user","content":"Reply only with OK."}],max_tokens=8,stream=False)) else "EMPTY RESPONSE"
    except Exception:return "FAILED"
def check_gemini_health():
    if not GEMINI_API_KEY:return "NOT CONFIGURED"
    try:return "ONLINE" if gemini_call("Reply only with OK.",system="You are a health check.") else "EMPTY RESPONSE"
    except Exception:return "FAILED"
def check_groq_models():
    if not client:return "NOT CONFIGURED"
    try:
        ids={getattr(m,"id","") for m in client.models.list().data}; missing=[x for ok,x in ((CHAT_MODEL in ids,"chat"),(VISION_MODEL in ids,"vision")) if not ok]; return "AVAILABLE" if not missing else "MISSING "+", ".join(missing)
    except Exception:return "UNKNOWN"

# ============================================================
# WEBHOOK
# ============================================================
_seen_ids=OrderedDict(); _seen_lock=threading.Lock(); SEEN_MAX=2000
def verify_signature(raw_body,header):
    if not APP_SECRET:return True
    if not header or not header.startswith("sha256="):return False
    return hmac.compare_digest(hmac.new(APP_SECRET.encode(),raw_body,hashlib.sha256).hexdigest(),header[7:])
def is_duplicate(msg_id):
    if not msg_id:return False
    with _seen_lock:
        if msg_id in _seen_ids:return True
        _seen_ids[msg_id]=time.time()
        while len(_seen_ids)>SEEN_MAX:_seen_ids.popitem(last=False)
    return False
if not APP_SECRET:print("[WARNING] WHATSAPP_APP_SECRET is not set. Webhook signatures are NOT verified.")

# ============================================================
# LIVE WEB SEARCH (Gemini Google Search grounding, uses GEMINI_API_KEY)
# ============================================================
LIVE_ENABLED=bool(GEMINI_API_KEY)  # live search runs on your Gemini key (Google Search grounding)
_LIVE_RE=re.compile(r"\b(today|tonight|yesterday|tomorrow|latest|currently|current|right now|news|breaking|headlines?|price of|stock|share price|exchange rate|score|scores|fixtures?|weather|forecast|trending|this (?:week|month|year)|just (?:happened|announced|released)|update on|release date|who (?:is|are) the (?:current )?(?:president|prime minister|ceo|governor|king|champion)|2026)\b",re.I)
def needs_live(text):return bool(_LIVE_RE.search(text or ""))
def _now_str():
    try:return datetime.now(_tz()).strftime("%A, %d %B %Y, %I:%M %p %Z")
    except Exception:return datetime.utcnow().strftime("%A, %d %B %Y, %H:%M UTC")
_LIVE_SYSTEM="You answer using live web information for WhatsApp. Be concise. Give dates for time-sensitive facts. If sources disagree or the evidence is thin, say so plainly. Never invent facts, quotes or sources."
def gemini_search_call(prompt,system=_LIVE_SYSTEM,timeout=45):
    payload={"system_instruction":{"parts":[{"text":system}]},"contents":[{"role":"user","parts":[{"text":prompt}]}],"tools":[{"google_search":{}}],"generationConfig":{"maxOutputTokens":1200}}
    r=requests.post(f"{GEMINI_API_URL}{GEMINI_MODEL}:generateContent",headers={"Content-Type":"application/json","x-goog-api-key":GEMINI_API_KEY},json=payload,timeout=timeout); r.raise_for_status(); data=r.json()
    text=_extract_gemini_text(data)
    if not text:raise ValueError("Gemini returned an empty grounded response")
    names=[]
    try:
        for ch in ((data.get("candidates") or [{}])[0].get("groundingMetadata") or {}).get("groundingChunks",[]):
            t=((ch.get("web") or {}).get("title") or "").strip()
            if t and t not in names:names.append(t)
    except Exception:pass
    return text,names[:3]
def _fmt_sources(names):return ("\n\n🌐 Sources: "+", ".join(names)) if names else "\n\n🌐 Live web search"
def live_answer(question,query=None):
    errors=[]
    if GEMINI_API_KEY:
        try:
            text,names=gemini_search_call(f"Current date/time: {_now_str()}\n\n{question}"); return text.strip()+_fmt_sources(names)
        except Exception as e:errors.append(f"Gemini: {e!r}"); print("[LIVE GEMINI ERROR]",repr(e))
    raise RuntimeError("; ".join(errors) or "Live search is not configured")
def smart_ask(prompt,user,force=False,query=None):
    if LIVE_ENABLED and (force or needs_live(prompt)):
        if check_rate_limit(user,"live"):
            try:
                ctx=build_memory_context(user,prompt,skip_last_user=True,max_chars=1200)
                return live_answer((ctx+"\n\n" if ctx else "")+f"User question: {prompt}",query or prompt)
            except Exception as e:print("[LIVE FALLBACK]",repr(e))
        elif force:return "You've reached the hourly live-search limit. Try again later."
    return ai_call(prompt,user)
def _verify_claims(user,media_id,mime):
    if not LIVE_ENABLED:return "\n\n🌐 Live fact-check is off (add GEMINI_API_KEY)."
    if not check_rate_limit(user,"live"):return "\n\n🌐 Live fact-check skipped: hourly live-search limit reached."
    try:
        claims=vision_call(media_id,"List any specific factual claims, headlines, quotes, dates, statistics or event descriptions shown in this image as plain text (max 4 short lines). If the image contains no checkable claims, reply exactly NONE.",mime).strip()
        if not claims or claims.upper().startswith("NONE"):return "\n\n🌐 No checkable text claims found in the image."
        res=live_answer(f"Fact-check these claims taken from an image, using current reliable sources. For each claim say Supported, Disputed, False or Unverified, with one short reason and the date of the evidence. Do not guess.\n\nClaims:\n{claims}",claims[:300])
        return "\n\n〔 *LIVE FACT-CHECK* 〕\n"+res
    except Exception as e:print("[VERIFY LIVE ERROR]",repr(e)); return "\n\n🌐 Live fact-check failed. Try `.verify <claim>` instead."

_CONTROL_PREFIXES=(".remind",".health",".status",".menu",".about",".help",".voice",".users",".ban",".unban",".image",".describe",".verify",".solve",".math",".read",".pint",".photo",".wallpaper",".anime",".comic",".edu",".pinterest",".play",".create",".imagine",".doc")
def _is_control_command(tl):return tl.startswith(_CONTROL_PREFIXES)

def _send_vision_for_user(user,command):
    saved=user_waiting_image.get(user)
    if saved and time.time()-saved.get("saved_at",0)>IMAGE_WAIT_TIMEOUT:user_waiting_image.pop(user,None);saved=None
    if not saved:return send_text(user,"No recent image is waiting. Send an image first.")
    u=saved.get("media_id"); mime=saved.get("mime_type","image/jpeg")
    if command==".solve":r=solve_image_problem(u,mime)
    elif command==".math":r=solve_image_problem(u,mime,"math")
    elif command==".read":r=vision_call(u,"Extract readable text only. Preserve useful line breaks and mark unclear text as [unclear].",mime)
    elif command==".verify":r=vision_call(u,"Analyze visible indicators of AI generation, manipulation, editing or misleading presentation. State evidence and limitations; do not claim forensic certainty.",mime); r+=_verify_claims(user,u,mime)
    else:r=vision_call(u,"Describe this image naturally for WhatsApp. Mention setting, objects, actions, colors, composition and readable text only when legible. Do not guess identities.",mime)
    add_to_memory(user,"assistant",r);send_text(user,r)

@app.route("/webhook",methods=["GET","POST"])
def webhook():
    if request.method=="GET":
        challenge=request.args.get("hub.challenge"); token=request.args.get("hub.verify_token")
        if VERIFY_TOKEN and token!=VERIFY_TOKEN:return "Forbidden",403
        return challenge or "OK",200
    raw=request.get_data()
    if not verify_signature(raw,request.headers.get("X-Hub-Signature-256")):return "Forbidden",403
    data=request.get_json(silent=True); threading.Thread(target=process_webhook,args=(data,),daemon=True).start(); return "OK",200

def process_webhook(data):
    try:
        msg=None
        for entry in (data or {}).get("entry",[]):
            for change in entry.get("changes",[]):
                for candidate in (change.get("value") or {}).get("messages",[]):
                    if candidate.get("from"):msg=candidate;break
                if msg:break
            if msg:break
        if not msg:return "OK",200
        user=msg.get("from");
        if not user or is_duplicate(msg.get("id")):return "OK",200
        text=(msg.get("text") or {}).get("body","").strip(); tl=text.lower()
        ok,auth=check_auth(user,text)
        if not ok:send_text(user,auth);return "OK",200
        if auth:send_text(user,auth);send_text(user,get_menu());return "OK",200
        if is_explicit_request(text):send_text(user,safety_block_message());return "OK",200
        if tl.startswith(".ban ") or tl.startswith(".unban "):
            if user!=OWNER_NUMBER:return send_text(user,"Owner only command.") or "OK"
            target=text.split(None,1)[1].strip()
            if tl.startswith(".ban "):
                if target==OWNER_NUMBER:return send_text(user,"Owner protection: that number cannot be banned.") or "OK"
                BANNED_USERS.add(target);authenticated_users.discard(target);save_memory();send_text(user,f"Banned: {target}")
            else:BANNED_USERS.discard(target);save_memory();send_text(user,f"Unbanned: {target}")
            return "OK",200
        if tl==".users":send_text(user,get_users_list() if user==OWNER_NUMBER else "Owner only command.");return "OK",200

        if msg.get("type")=="image":
            d=msg.get("image") or {}; mid=d.get("id")
            if not mid:return send_text(user,"I couldn't retrieve that image.") or "OK"
            user_waiting_image[user]={"media_id":mid,"mime_type":d.get("mime_type") or "image/jpeg","saved_at":time.time()};send_text(user,"Image received.\n\nSend `.describe`, `.verify`, `.solve`, `.math`, or `.read`.");return "OK",200
        if msg.get("type")=="document":
            d=msg.get("document") or {}; mid=d.get("id"); filename=d.get("filename") or "document.txt"
            try:
                r=requests.get(graph_url(mid),headers={"Authorization":f"Bearer {WHATSAPP_TOKEN}"},timeout=20);r.raise_for_status();info=r.json();url=info.get("url");mime=info.get("mime_type") or "application/octet-stream";rawdoc=requests.get(url,headers={"Authorization":f"Bearer {WHATSAPP_TOKEN}"},timeout=30);rawdoc.raise_for_status();textdoc=parse_document(filename,rawdoc.content,mime);note=f"\n\nNote: this document is long, so I will use the first {DOCUMENT_CONTEXT_LIMIT:,} characters." if len(textdoc)>DOCUMENT_CONTEXT_LIMIT else "";store_document_context(user,filename,textdoc);send_text(user,f"Document received: *{filename}*\n\nI extracted the readable text. Ask me anything about it, or use `.doc <question>`.{note}")
            except Exception as e:print("[DOCUMENT ERROR]",repr(e));send_text(user,f"I couldn't read that document.\n\n{str(e)[:500]}")
            return "OK",200
        if msg.get("type")=="audio":
            routed=False
            try:
                send_text(user,"🎤 Listening...");a,m=download_whatsapp_audio((msg.get("audio") or {}).get("id"));tr=transcribe_audio(a,m)
                if not tr:return send_text(user,"I couldn't make out what you said. Try again.") or "OK"
                if is_explicit_request(tr):send_text(user,safety_block_message());return "OK",200
                spoken=re.sub(r"\b([ap])\.m\.?",r"\1m",tr,flags=re.I).strip()
                if re.match(r"^\.?remind\s+me\b",spoken,re.I):
                    text="."+spoken.lstrip(".").rstrip(" .!?");tl=text.lower();routed=True;send_text(user,f"🎤 {tr}")
                else:
                    add_to_memory(user,"user",f"[Voice] {tr}");result=smart_ask(tr,user);add_to_memory(user,"assistant",result)
                    if VOICE_MODE in {"voice","audio","on"} or user in voice_enabled_users:
                        if not send_voice_reply(user,result):send_text(user,result)
                    else:send_text(user,f"🎤 {tr}\n\n{result}")
            except Exception as e:print("[VOICE INPUT ERROR]",repr(e));send_text(user,"I couldn't process that voice note. Check `.health` or try again.")
            if not routed:return "OK",200

        if tl in {".voice",".voice status"}:send_text(user,f"Voice mode: *{VOICE_MODE}*\n\nUse `.voice on` or `.voice off`.");return "OK",200
        if tl==".voice on":voice_enabled_users.add(user);send_text(user,"Voice replies enabled for you.");return "OK",200
        if tl==".voice off":voice_enabled_users.discard(user);send_text(user,"Voice replies disabled for you.");return "OK",200
        if text and not _is_control_command(tl):add_to_memory(user,"user",text[5:].strip() if tl.startswith(".ask ") else text)
        learned=learn_fact(user,text)
        if learned:add_to_memory(user,"assistant",learned);send_text(user,learned);return "OK",200
        if tl=="forget me":forget_user_memory(user);send_text(user,"Memory cleared.");return "OK",200

        if tl.startswith(".remind"):
            if tl in {".remind",".remind list",".reminders"}:
                active=list_reminders(user)
                if not active:send_text(user,"You have no active reminders.\n\nTry `.remind me to study at 8pm`.")
                else:send_text(user,"〔 *REMINDERS* 〕\n\n"+"\n".join(f"• `{r['id']}` — {r['task']} — {_fmt_due(r)}" for r in active)+"\n\nCancel with `.remind cancel <id>`")
                return "OK",200
            if tl.startswith(".remind cancel"):
                parts=text.split(None,2)
                if len(parts)<3:send_text(user,"Usage: `.remind cancel <id>`")
                elif cancel_reminder(user,parts[2].strip().strip("`")):send_text(user,f"Reminder `{parts[2].strip().strip('`')}` cancelled.")
                else:send_text(user,"Reminder not found.")
                return "OK",200
            task,when=parse_reminder(text)
            if not task or not when:return send_text(user,"I couldn't understand that reminder (or the time has already passed).\n\nTry:\n`.remind me to study at 8pm`\n`.remind me tomorrow at 7am to call John`\n`.remind me in 30 minutes to check the oven`\n`.remind me on friday at 5pm to pay rent`\n`.remind me to pray at noon`") or "OK"
            if user not in ({OWNER_NUMBER}|ADMIN_NUMBERS) and len(list_reminders(user))>=MAX_REMINDERS_PER_USER:return send_text(user,f"You already have {MAX_REMINDERS_PER_USER} active reminders. Cancel one with `.remind cancel <id>`.") or "OK"
            rid=create_reminder(user,task,when);send_text(user,f"⏰ Reminder set.\n\n*{task}*\n{when.astimezone(_tz()).strftime('%d %b %Y, %I:%M %p')}\nID: `{rid}`");return "OK",200

        if tl==".health":
            health=f"〔 *ARIA HEALTH* 〕\n\n*Version:* {VERSION}\n*Runtime:* {get_runtime()}\n\n〔 *CORE* 〕\n*WhatsApp:* {'CONFIGURED' if WHATSAPP_TOKEN and PHONE_NUMBER_ID else 'NOT CONFIGURED'}\n*Memory:* OK\n\n〔 *GROQ* 〕\n*Chat:* {check_groq_chat_health()}\n*Models:* {check_groq_models()}\n*Vision:* {'READY' if client else 'NOT CONFIGURED'}\n*Chat Model:* {CHAT_MODEL}\n*Vision Model:* {VISION_MODEL}\n\n〔 *GEMINI* 〕\n*API:* {check_gemini_health()}\n*Vision:* {'READY' if GEMINI_API_KEY else 'NOT CONFIGURED'}\n\n〔 *VOICE* 〕\n*Mode:* {VOICE_MODE}\n*STT:* {STT_MODEL}\n*TTS:* {TTS_MODEL} / {TTS_VOICE}\n\n〔 *SECURITY* 〕\n*Owner:* PROTECTED\n*Password:* {'ENVIRONMENT' if ARIA_PASSWORD else 'DISABLED'}\n*Content Safety:* ENABLED";send_text(user,health);return "OK",200

        if tl.startswith(".verify "):
            claim=text[8:].strip()[:1000]
            if not LIVE_ENABLED:r="Live verification isn't configured. Add GEMINI_API_KEY."
            elif not check_rate_limit(user,"live"):r="You've reached the hourly live-search limit. Try again later."
            else:
                send_text(user,"🌐 Checking live sources...")
                try:r="〔 *LIVE VERIFY* 〕\n\n"+live_answer(f"Fact-check this claim using current reliable sources. Start with a verdict (Supported / Disputed / False / Unverified), then give 2-4 short reasons with dates. Do not guess.\n\nClaim: {claim}",claim[:300])
                except Exception as e:print("[VERIFY ERROR]",repr(e));r="Live check failed. Try again shortly."
            add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200
        if tl in {".search",".news"}:send_text(user,"Usage: `.search <question>` or `.news <topic>`");return "OK",200
        if tl.startswith((".search ",".news ")):
            cmd,_,q=text.partition(" ");q=q.strip()[:500]
            if not LIVE_ENABLED:send_text(user,"Live search isn't configured. Add GEMINI_API_KEY.");return "OK",200
            send_text(user,"🌐 Searching...")
            r=smart_ask(f"What is the latest news on: {q}? Give 3-5 short bullet points, each with its date." if cmd.lower()==".news" else q,user,force=True,query=q)
            add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200

        if tl==".image" or tl.startswith(".image "):
            send_text(user,"`.image` is not a vision follow-up in ARIA v15.2.\n\nFor AI image generation use:\n`.create <prompt>`\n\nFor analysing a WhatsApp image, send the image first, then use `.describe`, `.read`, `.solve`, `.math`, or `.verify`.");return "OK",200
        if tl.startswith(".describe ask "):
            saved=user_waiting_image.get(user)
            if not saved:return send_text(user,"No recent image is waiting. Send an image first.") or "OK"
            q=text[len(".describe ask "):].strip();r=vision_call(saved["media_id"],f"Answer this follow-up question about the image. Do not invent details.\n\nQuestion: {q}",saved.get("mime_type"));add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200
        if tl in {".describe",".describe detailed",".verify",".solve",".math",".read"}:
            saved=user_waiting_image.get(user)
            if saved and time.time()-saved.get("saved_at",0)>IMAGE_WAIT_TIMEOUT:user_waiting_image.pop(user,None);saved=None
            if not saved:return send_text(user,"No recent image is waiting. Send an image first.") or "OK"
            cmd=tl
            if cmd==".describe detailed":r=vision_call(saved["media_id"],"Describe this image in detail, including composition, setting, colors, objects, actions, style and readable text. Do not guess.",saved.get("mime_type"))
            else:
                _send_vision_for_user(user,cmd);return "OK",200
            add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200

        for command,handler in [(".photo",pint1_unsplash),(".wallpaper",pint2_pexels),(".anime",pint3_anime),(".comic",pint4_comics),(".edu",pint5_education),(".pinterest",pint_pinterest)]:
            if tl.startswith(command+" "):handler(user,text[len(command):].strip());return "OK",200
        if tl==".pint":send_text(user,"〔 PINT 〕\n\nChoose `.photo`, `.wallpaper`, `.anime`, `.comic`, `.edu`, or `.pinterest`.");return "OK",200
        for i,handler in enumerate([pint1_unsplash,pint2_pexels,pint3_anime,pint4_comics,pint5_education,pint_pinterest],1):
            if tl.startswith(f".pint{i} "):handler(user,text[len(f".pint{i}"):].strip());return "OK",200

        if tl==".menu":send_text(user,get_menu());return "OK",200
        if tl==".about":send_text(user,get_about());return "OK",200
        if tl==".status":send_text(user,get_status());return "OK",200
        if tl.startswith(".help"):
            c=text[5:].strip().lower(); helpmap={".solve":"Send an image, then `.solve`.",".math":"Send an image, then `.math`.",".read":"Send an image, then `.read`.",".describe":"Send an image, then `.describe`; use `.describe ask <question>` for a question about that image.",".verify":"Send an image, then `.verify` (checks edits/AI signs and fact-checks text claims live), or use `.verify <claim>`.",".create":"Generate an image with `.create <prompt>`.",".remind":"Set reminders with `.remind me to ... at 8pm`, `.remind me tomorrow at 7am to ...`, or `.remind me in 30 minutes to ...`.",".document":"Send a supported document and ask questions about it."};send_text(user,"Usage: `.help <command>`\n\n"+(helpmap.get(c,"Try `.menu`.") if c else "Try `.help solve`, `.help create`, `.help remind`, or `.help document`."));return "OK",200
        if tl.startswith(".play"):
            q=text[5:].strip();send_text(user,get_youtube_link(q) if q else "Usage: `.play <song name>`");return "OK",200
        if tl.startswith(".create ") or tl.startswith(".imagine "):
            prompt=text.split(None,1)[1].strip()[:500]
            if not check_rate_limit(user,"create"):send_text(user,"You've reached the hourly image limit. Try again later.");return "OK",200
            imagine_generate(user,prompt);return "OK",200
        if tl in {".create",".imagine"}:send_text(user,"Usage: `.create <prompt>`\n\nExample: `.create cyberpunk city at night`");return "OK",200
        if tl.startswith(".explain") or re.match(r"^explain\s+.+\s+[1-6]$",tl):
            ct=text[1:] if text.startswith(".") else text; sp=ct.split(None,1); args=sp[1].strip() if len(sp)>1 else ""
            mm=re.match(r"^(.+?)\s+([1-6])$",args)
            topic,field=(mm.group(1).strip(),mm.group(2)) if mm else (args,"all")
            r="Usage: `.explain <topic>`" if not topic else ai_explain(topic,field,user);add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200
        if tl.startswith(".ask "):
            r=smart_ask(text[5:].strip(),user);add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200
        if tl==".ask":send_text(user,"Usage: `.ask <question>`");return "OK",200
        if tl.startswith(".summarize "):r=ai_call(f"Summarize this clearly and briefly:\n\n{text[11:].strip()}",user);add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200
        if tl==".summarize":send_text(user,"Usage: `.summarize <text>`");return "OK",200
        if tl.startswith(".translate "):
            p=text.split(" ",2)
            if len(p)<3:send_text(user,"Usage: `.translate <language> <text>`")
            else:r=ai_call(f"Translate into {p[1]}. Return only the natural translation.\n\n{p[2]}",user);add_to_memory(user,"assistant",r);send_text(user,r)
            return "OK",200
        if tl.startswith(".define "):r=ai_call(f"Define '{text[8:].strip()}'. Give a concise definition and one short example.",user);add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200
        if tl.startswith(".study "):r=ai_call(f"Create concise study notes for '{text[7:].strip()}'. Include definition, key ideas, one example, common mistake, and 3 exam-focused points.",user);add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200
        if tl.startswith(".quiz "):r=ai_call(f"Create a 5-question quiz on '{text[6:].strip()}'. Do not reveal answers yet; ask the user to reply with their answers.",user);add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200
        if re.match(r"^\.(?:doc|document)(?:\s|$)",tl) or (get_document_context(user) and tl.startswith(("document ","ask document"))):
            q=re.sub(r"^\.?(?:ask\s+)?(?:document|doc)(?:\s+|$)","",text,flags=re.I).strip() or "Summarize the document."
            r=ask_about_document(user,q);add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200
        if get_document_context(user) and tl.startswith(("what ","who ","when ","where ","why ","how ","summarize", "explain")):
            # Only use document context when it is clearly the active context.
            r=ask_about_document(user,text);add_to_memory(user,"assistant",r);send_text(user,r);return "OK",200
        if not text:send_text(user,"Send me a message.");return "OK",200
        r=smart_ask(text,user);add_to_memory(user,"assistant",r);send_text(user,r)
    except Exception as e:print("[WEBHOOK ERROR]",repr(e))
    return "OK",200

@app.route("/")
def home():return f"ARIA {VERSION} Running | Graph API {GRAPH_API_VERSION} | Chat {CHAT_MODEL} | Vision {VISION_MODEL} | Gemini {GEMINI_MODEL}"

@app.route("/ping")
def ping():return "pong",200

if __name__=="__main__":
    port=int(os.getenv("PORT","5000"));app.run(host="0.0.0.0",port=port)










